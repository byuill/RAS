"""Build a navigable, editable Word document with deterministic page layout."""
from pathlib import Path
from datetime import datetime, timezone
import json
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree

from guide_paths import ROOT
NAVY='17354C';TEAL='287C83';GRAY='626D76';LIGHT='EDF3F6'
PARTS=[
    (1,7,'I  Prepare a defensible comparison'),
    (8,12,'II  Trace the volume calculation'),
    (13,18,'III  Diagnose model agreement'),
    (19,24,'IV  Test assumptions and uncertainty'),
    (25,29,'V  Extend and document the review'),
    (30,34,'VI  Reference'),
]

def shade(cell, fill):
    props=cell._tc.get_or_add_tcPr()
    el=OxmlElement('w:shd');el.set(qn('w:fill'),fill);props.append(el)

def no_split(row):
    el=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(el)

def field(paragraph, code):
    run=paragraph.add_run()
    for tag,value in [('begin',None),(None,code),('end',None)]:
        el=OxmlElement('w:instrText' if tag is None else 'w:fldChar')
        if tag is None:el.text=value;el.set(qn('xml:space'),'preserve')
        else:el.set(qn('w:fldCharType'),tag)
        run._r.append(el)

def internal_link(paragraph, text, anchor):
    el=OxmlElement('w:hyperlink');el.set(qn('w:anchor'),anchor)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),NAVY);props.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'18');props.append(size)
    run.append(props);txt=OxmlElement('w:t');txt.text=text;run.append(txt)
    el.append(run);paragraph._p.append(el)

class Guide:
    def __init__(self):
        self.doc=Document()
        self.pages=[];self.figure_count=0
        sec=self.doc.sections[0]
        sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65)
        sec.left_margin=Inches(.78);sec.right_margin=Inches(.72)
        sec.header_distance=Inches(.28);sec.footer_distance=Inches(.29)
        sec.different_first_page_header_footer=True
        styles=self.doc.styles
        normal=styles['Normal']
        normal.font.name='Calibri';normal.font.size=Pt(10.5)
        normal.font.color.rgb=RGBColor.from_string(NAVY)
        normal.paragraph_format.line_spacing=1.10
        normal.paragraph_format.space_after=Pt(6)
        normal.paragraph_format.widow_control=True
        for name,size,color in [('Title',38,NAVY),('Subtitle',17,GRAY),
            ('Heading 1',21,NAVY),('Heading 2',12.3,TEAL),('Heading 3',10.5,NAVY)]:
            styles[name].font.name='Calibri';styles[name].font.size=Pt(size)
            styles[name].font.color.rgb=RGBColor.from_string(color)
            styles[name].paragraph_format.space_after=Pt(8 if name=='Heading 1' else 5)
            styles[name].paragraph_format.space_before=Pt(0 if name in ('Title','Heading 1') else 6)
            styles[name].paragraph_format.keep_with_next=True
        styles['Caption'].font.name='Calibri';styles['Caption'].font.size=Pt(9)
        styles['Caption'].font.color.rgb=RGBColor.from_string(GRAY)
        styles['Caption'].paragraph_format.space_after=Pt(8)
        for name in ['List Bullet','List Number']:
            styles[name].font.name='Calibri';styles[name].font.size=Pt(10.5)
            styles[name].paragraph_format.space_after=Pt(5)
            styles[name].paragraph_format.line_spacing=1.10
        code=styles.add_style('Guide Code',1)
        code.font.name='Liberation Mono';code.font.size=Pt(9)
        code.font.color.rgb=RGBColor.from_string(NAVY)
        code.paragraph_format.space_after=Pt(5);code.paragraph_format.line_spacing=1.05
        # Keep code monospaced in renderers and in Office font substitution.
        font=OxmlElement('w:font');font.set(qn('w:name'),'Liberation Mono')
        for tag,value in [('altName','Courier New'),('family','modern'),('pitch','fixed')]:
            node=OxmlElement('w:'+tag);node.set(qn('w:val'),value);font.append(node)
        font_part=self.doc.part.related_parts[next(
            rel.rId for rel in self.doc.part.rels.values() if rel.reltype==RT.FONT_TABLE)]
        font_table=parse_xml(font_part.blob);font_table.append(font)
        font_part._blob=etree.tostring(font_table,encoding='UTF-8',xml_declaration=True,standalone=True)
        header=sec.header.paragraphs[0]
        header.text='BATHYMETRY CALIBRATION  /  USER GUIDE'
        header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
        footer=sec.footer.paragraphs[0]
        footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_ALIGN_PARAGRAPH.RIGHT)
        footer.add_run('Version 1.1  •  12 October 2026\t')
        field(footer,'PAGE')
        for run in footer.runs:
            run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string(GRAY)
        self.doc.core_properties.title='Bathymetry Calibration — User Guide for HEC-RAS Technical Review'
        self.doc.core_properties.subject='Workflow, bed-volume calculations, QA/QC, sensitivity, uncertainty and review reporting'
        self.doc.core_properties.author='RAS project documentation'
        self.doc.core_properties.created=datetime.now(timezone.utc)
        self.doc.core_properties.modified=datetime.now(timezone.utc)
        self.doc.core_properties.keywords='HEC-RAS, bathymetry, bed change, calibration, user guide, QA/QC'
        self.cover()
        self.doc.add_page_break()
        self.doc.add_heading('How to use this guide',level=1)
        self.p('For experienced HEC-RAS modelers and reviewers. The six parts follow the work from verified inputs '
               'to traceable calculations, diagnostic interpretation and review evidence. Section numbers are stable references.')
        self.p('**First run:** §§01, 03–07, 13, 27.  **Audit the curve:** §§07–12, 19–24.\n'
               '**Prepare a formal review:** §§02, 13–18, 26–29, 34. Click a contents entry to jump to it.')
        self.toc=self.doc.add_table(rows=1,cols=4)
        self.toc.alignment=WD_TABLE_ALIGNMENT.CENTER;self.toc.autofit=False
        for column,width in zip(self.toc.columns,[3.15,.35,3.15,.35]):
            column.width=Inches(width)
        for c,text,w in zip(self.toc.rows[0].cells,['Topic','p.','Topic','p.'],[3.15,.35,3.15,.35]):
            c.width=Inches(w);c.text=text;shade(c,NAVY)
            for run in c.paragraphs[0].runs:
                run.font.size=Pt(9);run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        self.toc_after=self.doc.add_paragraph()
        self.toc_after.paragraph_format.space_before=Pt(7)
        self.toc_after.add_run('Scope: byuill/RAS, bathymetry_calibration, sunday_improvements, commit e435355. '
            'The guide describes that implementation; it does not certify project survey data or model performance.').font.size=Pt(9)

    def cover(self):
        self.doc.add_paragraph('HEC-RAS  /  SEDIMENT MODEL REVIEW').runs[0].font.color.rgb=RGBColor.from_string(TEAL)
        self.doc.add_paragraph()
        self.doc.add_paragraph('Bathymetry\nCalibration',style='Title')
        self.doc.add_paragraph('User guide for modelers\nand technical reviewers',style='Subtitle')
        self.doc.add_paragraph()
        self.p('A practical guide to operating the workbench, understanding cumulative bed-volume curves, '
               'reviewing observed and modeled change, and reporting the strength and limits of the comparison.')
        self.callout('What you will be able to explain',
            'How four bed profiles become section areas, reach volumes and cumulative curves; '
            'why opposing erosion and deposition can cancel; how calculation assumptions affect bias; '
            'and what survey-error intervals do and do not establish.')
        self.doc.add_paragraph()
        self.p('**Displayed units:** feet • square feet • cubic yards • model river miles')
        self.p('**Tool:** RAS / bathymetry_calibration\n**Implementation:** sunday_improvements • e435355\n'
               '**Guide:** Version 1.1 • 12 October 2026')
        self.doc.add_paragraph()
        self.p('All worked examples and example model–survey plots are illustrative. '
               'They are not results for the Mississippi River project or evidence that a particular HEC-RAS model is calibrated.')
        self.doc.paragraphs[-1].runs[0].font.color.rgb=RGBColor.from_string(GRAY)

    def page(self,title):
        number=len(self.pages)+1
        part=next(p for p in PARTS if p[0]<=number<=p[1])
        if number==part[0]:
            label=self.doc.add_heading('PART '+part[2],level=1)
            label.paragraph_format.page_break_before=True
        else:
            label=self.doc.add_paragraph('PART '+part[2])
            label.paragraph_format.page_break_before=True
        label.paragraph_format.space_after=Pt(5)
        label.paragraph_format.keep_with_next=True
        for run in label.runs:
            run.font.size=Pt(8.5);run.font.bold=True
            run.font.color.rgb=RGBColor.from_string(TEAL)
        heading=self.doc.add_heading(f'{number:02d}  {title}',level=2)
        heading.paragraph_format.space_before=Pt(0)
        heading.paragraph_format.space_after=Pt(8)
        for run in heading.runs:
            run.font.size=Pt(21);run.font.color.rgb=RGBColor.from_string(NAVY)
        bookmark=f's{number:02d}'
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),bookmark)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
        heading._p.insert(0,start);heading._p.append(end)
        self.pages.append({'title':title,'bookmark':bookmark,'page':number+2})

    def p(self,text,style=None):
        paragraph=self.doc.add_paragraph(style=style)
        for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
            run=paragraph.add_run(part)
            if i%2:run.bold=True
        return paragraph

    def h(self,text):
        p=self.doc.add_heading(text,level=3)
        for run in p.runs:
            run.font.size=Pt(12.3);run.font.color.rgb=RGBColor.from_string(TEAL)
    def bullets(self,items):
        for item in items:self.p(item,'List Bullet')
    def steps(self,items):
        for i,item in enumerate(items,1):self.p(f'**{i}.** {item}')
    def code(self,text):self.p(text,'Guide Code')
    def equations(self,rows):
        table=self.table(['Calculation','Meaning / units'],rows,[4.55,2.45],font=10)
        for row in table.rows[1:]:
            p=row.cells[0].paragraphs[0]
            equation=p.text;p.clear()
            pos=0
            for match in re.finditer(r'([_^])\{(.*?)\}',equation):
                run=p.add_run(equation[pos:match.start()]);run.font.name='Cambria Math';run.font.size=Pt(11)
                run=p.add_run(match.group(2));run.font.name='Cambria Math';run.font.size=Pt(10)
                run.font.subscript=match.group(1)=='_';run.font.superscript=match.group(1)=='^'
                pos=match.end()
            run=p.add_run(equation[pos:]);run.font.name='Cambria Math';run.font.size=Pt(11)
    def callout(self,title,text):
        table=self.doc.add_table(rows=1,cols=1);table.alignment=WD_TABLE_ALIGNMENT.CENTER
        cell=table.cell(0,0);shade(cell,LIGHT);no_split(table.rows[0])
        p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4)
        r=p.add_run(title);r.bold=True;r.font.color.rgb=RGBColor.from_string(TEAL);r.font.size=Pt(10.5)
        p=cell.add_paragraph(text);p.paragraph_format.space_after=Pt(5)
        for r in p.runs:r.font.size=Pt(10)
        self.spacer()

    def spacer(self):
        p=self.doc.add_paragraph()
        p.paragraph_format.line_spacing=Pt(1)
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(3)

    def table(self,headers,rows,widths=None,font=9.2):
        table=self.doc.add_table(rows=1,cols=len(headers))
        table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        widths=widths or [7/len(headers)]*len(headers)
        for column,width in zip(table.columns,widths):
            column.width=Inches(width)
        for cell,text,w in zip(table.rows[0].cells,headers,widths):
            cell.width=Inches(w);cell.text=text;shade(cell,NAVY)
            for run in cell.paragraphs[0].runs:
                run.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(font)
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i,values in enumerate(rows):
            row=table.add_row();no_split(row)
            for cell,text,w in zip(row.cells,values,widths):
                cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text=str(text)
                if i%2==0:shade(cell,'F3F6F8')
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.line_spacing=1.03
                    for run in p.runs:run.font.size=Pt(font)
        self.spacer()
        return table

    def image(self,name,caption,width=6.85):
        self.figure_count+=1
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
        run=p.add_run();inline=run.add_picture(str(ROOT/'figures'/(name+'.png')),width=Inches(width))
        inline._inline.docPr.set('descr',caption)
        self.p(f'Figure {self.figure_count}. {caption}','Caption')

    def finish(self):
        def entries(parts):
            result=[]
            for first,last,title in parts:
                result.append({'part':title})
                result.extend(self.pages[first-1:last])
            return result
        left,right=entries(PARTS[:3]),entries(PARTS[3:])
        for i in range(max(len(left),len(right))):
            row=self.toc.add_row();no_split(row)
            items=[col[i] if i<len(col) else None for col in (left,right)]
            for offset,item in zip([0,2],items):
                if item:
                    if 'part' in item:
                        cell=row.cells[offset].merge(row.cells[offset+1])
                        cell.text=item['part'];shade(cell,LIGHT)
                        for run in cell.paragraphs[0].runs:
                            run.bold=True;run.font.color.rgb=RGBColor.from_string(TEAL)
                    else:
                        section=int(item['bookmark'][1:])
                        internal_link(row.cells[offset].paragraphs[0],f'{section:02d}  '+item['title'],item['bookmark'])
                        row.cells[offset+1].text=str(item['page'])
                for cell in row.cells[offset:offset+2]:
                    if i%2==0 and item and 'part' not in item:shade(cell,'F3F6F8')
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                        p.paragraph_format.line_spacing=1.03
                        for r in p.runs:r.font.size=Pt(9)
        settings=self.doc.settings.element
        update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
        filename=ROOT/'Bathymetry_Calibration_User_Guide.docx'
        self.doc.save(filename)
        (ROOT/'qa'/'expected_pages.json').write_text(json.dumps(self.pages,indent=2))
        words=sum(len(p.text.split()) for p in self.doc.paragraphs)
        words+=sum(len(c.text.split()) for t in self.doc.tables for r in t.rows for c in r.cells)
        print(f'{filename}: {len(self.pages)+2} planned pages, {self.figure_count} figures, approximately {words:,} words')

if __name__=='__main__':
    from guide_content import write_content
    guide=Guide();write_content(guide);guide.finish()
